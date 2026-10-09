"""Normalize Citywide incidents from the source workbook to CSV and JSON.

This script preserves source fields and deliberately leaves unsupported road IDs,
coordinates, road-level rainfall, and thresholds null. A small, explicit audit
registry records article checks documented in docs/source-verification-log.md.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from datetime import date
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "data" / "raw" / "kanpur_gurugram_all_waterlogging_locations.xlsx"
DEFAULT_OUTPUT = ROOT / "data" / "processed"
INCIDENT_SHEET = "Citywide incidents"

FIELDS = [
    "record_id", "event_id", "source_record_id", "road_id", "road_name",
    "location_name", "location_precision", "city", "latitude", "longitude",
    "event_date", "event_date_start", "event_date_end", "event_date_raw",
    "rainfall_amount_mm", "rainfall_scope", "rainfall_scope_period",
    "rainfall_is_road_specific", "rainfall_duration_minutes",
    "rainfall_duration_qualifier", "rainfall_duration_text",
    "rainfall_intensity_mm_hr", "severity", "severity_scope", "description",
    "source_name", "source_url", "evidence_status", "confidence",
    "threshold_status", "threshold_assessment", "threshold_limitations",
    "is_priority_location", "review_flags", "source_sheet", "source_excel_row",
    "source_city_raw", "source_location_raw", "source_event_date_raw",
    "source_rainfall_amount_raw", "source_rainfall_scope_raw",
    "source_duration_raw", "source_severity_context_raw",
    "source_evidence_note_raw", "source_priority_label_raw",
]

PUBLISHERS = {
    "timesofindia.indiatimes.com": "The Times of India",
    "www.hindustantimes.com": "Hindustan Times",
    "indianexpress.com": "The Indian Express",
    "www.indiatoday.in": "India Today",
    "www.dailypioneer.com": "The Daily Pioneer",
    "www.tribuneindia.com": "The Tribune",
}

THRESHOLD_LIMITATION = (
    "INSUFFICIENT EVIDENCE FOR THRESHOLD: rainfall and duration are contextual, "
    "not a verified road-specific causal pairing."
)

# Direct source checks documented in docs/source-verification-log.md (checked
# 2026-10-09). Status applies to the reported event/location claim, not to a
# road-specific rainfall threshold. Keep source workbook cells in source_* fields.
VERIFIED_SOURCE_CHECKS = {
    "HIST-LOC-0036": (
        "Direct article check (2026-10-09): confirms water accumulation at Chunniganj "
        "during the 2019-09-27 Kanpur event. The 89.4 mm figure is a CSA city gauge "
        "measurement, not a Chunniganj road measurement; no continuous duration or "
        "road-specific depth is reported."
    ),
    "HIST-LOC-0070": (
        "Direct article check (2026-10-09): confirms Govind Nagar was among Kanpur "
        "areas affected by waterlogging on 2025-07-12. The 35 mm figure is for the "
        "city over 24 hours; no road-specific rainfall or duration is reported."
    ),
    "HIST-LOC-0073": (
        "Direct article check (2026-10-09): confirms P Road was among Kanpur areas "
        "affected by waterlogging on 2025-07-12. The 35 mm figure is for the city "
        "over 24 hours; no road-specific rainfall or duration is reported."
    ),
    "HIST-LOC-0077": (
        "Direct article check (2026-10-09): confirms Gwaltoli was among Kanpur areas "
        "affected by waterlogging on 2025-07-12. The 35 mm figure is for the city "
        "over 24 hours; no road-specific rainfall or duration is reported."
    ),
    "HIST-LOC-0116": (
        "Direct article check (2026-10-09): confirms the Hero Honda Chowk underpass "
        "was closed and remained submerged after the 2018-08-28 waterlogging event. "
        "The article associates the event with 128 mm of rainfall but gives no "
        "continuous rainfall duration or underpass-specific gauge measurement."
    ),
    "HIST-LOC-0123": (
        "Direct article check (2026-10-09): confirms Rajeev Chowk was among Gurugram "
        "areas affected by waterlogging on 2022-05-23. The cited article gives no "
        "rainfall amount or duration. The workbook's approximately two-hour value "
        "was not confirmed and is therefore null in the normalized duration field."
    ),
    "HIST-LOC-0131": (
        "Direct article check (2026-10-09): confirms waterlogging at Golf Course "
        "Road during the 2022-08-07 Gurugram event. The report says rain lasted over "
        "two hours and gives 19 mm for Gurgaon by 5 pm; the rainfall figure is "
        "city/event context, not a road measurement. Knee-high depth applied only "
        "to some low-lying points, not specifically this road."
    ),
    "HIST-LOC-0150": (
        "Direct article check (2026-10-09): confirms waterlogging at Sheetla Mata "
        "Road during the 2022-08-07 Gurugram event. The report says rain lasted over "
        "two hours and gives 19 mm for Gurgaon by 5 pm; the rainfall figure is "
        "city/event context, not a road measurement. Knee-high depth applied only "
        "to some low-lying points, not specifically this road."
    ),
    "HIST-LOC-0155": (
        "Direct article check (2026-10-09): confirms waterlogging at Rajeev Chowk "
        "during the 2023-07-04 Gurugram event. The 65 mm figure is the city total "
        "reported by 4:30 pm; article describes heavy rain in a short span but "
        "states no exact duration or road-specific rainfall."
    ),
    "HIST-LOC-0161": (
        "Direct article check (2026-10-09): confirms waterlogging at Hero Honda "
        "Chowk during the 2023-07-04 Gurugram event. The 65 mm figure is the city "
        "total reported by 4:30 pm; article describes heavy rain in a short span "
        "but states no exact duration or road-specific rainfall."
    ),
}


def source_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def numeric_or_none(value: object) -> int | float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    return None


def parsed_dates(raw: str) -> tuple[str | None, str | None, str | None]:
    found = re.findall(r"\b\d{4}-\d{2}-\d{2}\b", raw)
    if len(found) == 1 and raw.strip() == found[0]:
        return found[0], None, None
    if len(found) == 2:
        return None, found[0], found[1]
    return None, None, None


def normalize_duration(raw: str) -> tuple[int | None, str]:
    """Convert only an explicit rain-duration statement; retain its qualifier."""
    text = raw.casefold()
    if re.search(r"about\s+2\.5\s+hours?", text):
        return 150, "APPROXIMATE"
    if re.search(r"over\s+2\s+hours?", text):
        return 120, "AT_LEAST"
    if (
        re.search(r"about\s+2\s+hours?", text)
        and "cleared" not in text
        and "reced" not in text
        and "traffic disrupted" not in text
    ):
        return 120, "APPROXIMATE"
    return None, "UNKNOWN"


def scope_type(raw: str) -> str:
    text = raw.casefold()
    if "district" in text:
        return "DISTRICT_TOTAL"
    if "gauge" in text or "observatory" in text:
        return "WEATHER_STATION"
    if "city" in text:
        return "CITY_TOTAL"
    return "UNKNOWN"


def source_name(url: str) -> str | None:
    host = urlparse(url).hostname
    if not host:
        return None
    return PUBLISHERS.get(host.casefold(), host)


def stable_event_id(city: str, date_raw: str, url: str) -> str:
    key = "|".join((city.strip().casefold(), date_raw.strip(), url.strip()))
    return "EVT-" + hashlib.sha256(key.encode("utf-8")).hexdigest()[:12]


def load_records(workbook_path: Path) -> tuple[list[dict], dict]:
    book = load_workbook(workbook_path, read_only=True, data_only=True)
    expected = {"Start here", "Citywide incidents", "Rainfall context", "Priority locations"}
    if set(book.sheetnames) != expected:
        raise ValueError(f"Unexpected worksheets: {book.sheetnames}")

    ws = book[INCIDENT_SHEET]
    headers = [source_text(cell.value) for cell in ws[5]]
    required_headers = [
        "City", "Reported location / road", "Event date", "Rainfall (mm)",
        "Rainfall scope / period", "Rain / water duration",
        "Reported severity (event context)", "Evidence note", "Priority location?", "Source link",
    ]
    if headers[:10] != required_headers:
        raise ValueError(f"Unexpected incident headers: {headers}")
    col = {name: index for index, name in enumerate(headers)}
    records: list[dict] = []

    for excel_row, cells in enumerate(ws.iter_rows(min_row=6, values_only=True), start=6):
        if not any(value is not None and value != "" for value in cells):
            continue
        city_raw = source_text(cells[col["City"]])
        location_raw = source_text(cells[col["Reported location / road"]])
        event_date_raw = source_text(cells[col["Event date"]])
        rainfall_scope_raw = source_text(cells[col["Rainfall scope / period"]])
        duration_raw = source_text(cells[col["Rain / water duration"]])
        severity_raw = source_text(cells[col["Reported severity (event context)"]])
        evidence_raw = source_text(cells[col["Evidence note"]])
        priority_raw = source_text(cells[col["Priority location?"]])
        url = source_text(cells[col["Source link"]])
        source_rainfall = cells[col["Rainfall (mm)"]]

        if not city_raw or not event_date_raw or not url:
            raise ValueError(f"Missing city, source date, or source URL at {INCIDENT_SHEET}!{excel_row}")

        event_date, date_start, date_end = parsed_dates(event_date_raw)
        duration, duration_qualifier = normalize_duration(duration_raw)
        location_name = location_raw or None
        location_precision = "NOT_ASSESSED"
        if "specific route unknown" in location_raw.casefold():
            location_name = re.sub(r"\s*\(specific route UNKNOWN\)", "", location_raw, flags=re.I).strip()
            location_precision = "AREA_ONLY"
        if location_raw.casefold() == "unknown":
            location_name = None
            location_precision = "UNKNOWN"

        rainfall_amount = numeric_or_none(source_rainfall)
        row_id = f"HIST-LOC-{excel_row:04d}"
        event_id = stable_event_id(city_raw, event_date_raw, url)
        records.append({
            "record_id": row_id,
            "event_id": event_id,
            "source_record_id": f"{INCIDENT_SHEET}!A{excel_row}:J{excel_row}",
            "road_id": None,
            "road_name": None,
            "location_name": location_name,
            "location_precision": location_precision,
            "city": city_raw,
            "latitude": None,
            "longitude": None,
            "event_date": event_date,
            "event_date_start": date_start,
            "event_date_end": date_end,
            "event_date_raw": event_date_raw,
            "rainfall_amount_mm": rainfall_amount,
            "rainfall_scope": scope_type(rainfall_scope_raw),
            "rainfall_scope_period": rainfall_scope_raw,
            "rainfall_is_road_specific": False,
            "rainfall_duration_minutes": duration,
            "rainfall_duration_qualifier": duration_qualifier,
            "rainfall_duration_text": duration_raw,
            "rainfall_intensity_mm_hr": None,
            "severity": severity_raw or None,
            "severity_scope": "EVENT_CONTEXT" if severity_raw else "UNKNOWN",
            "description": evidence_raw or None,
            "source_name": source_name(url),
            "source_url": url,
            "evidence_status": "REQUIRES_REVIEW",
            "confidence": None,
            "threshold_status": "UNKNOWN",
            "threshold_assessment": "INSUFFICIENT EVIDENCE FOR THRESHOLD",
            "threshold_limitations": THRESHOLD_LIMITATION,
            "is_priority_location": priority_raw.casefold().startswith("yes"),
            "review_flags": [],
            "source_sheet": INCIDENT_SHEET,
            "source_excel_row": excel_row,
            "source_city_raw": city_raw,
            "source_location_raw": location_raw,
            "source_event_date_raw": event_date_raw,
            "source_rainfall_amount_raw": numeric_or_none(source_rainfall) if rainfall_amount is not None else source_text(source_rainfall),
            "source_rainfall_scope_raw": rainfall_scope_raw,
            "source_duration_raw": duration_raw,
            "source_severity_context_raw": severity_raw,
            "source_evidence_note_raw": evidence_raw,
            "source_priority_label_raw": priority_raw,
        })

    for record in records:
        source_check = VERIFIED_SOURCE_CHECKS.get(record["record_id"])
        if source_check:
            record["evidence_status"] = "VERIFIED"
            record["description"] = source_check

    # The workbook's approximate duration for this source/event is not
    # supported by the cited India Today article. Clear it for every row linked
    # to that article, while retaining the original transcription in
    # source_duration_raw for auditability.
    may_2022_url = next(
        r["source_url"] for r in records if r["record_id"] == "HIST-LOC-0123"
    )
    for record in records:
        if record["source_url"] == may_2022_url:
            record["rainfall_duration_minutes"] = None
            record["rainfall_duration_qualifier"] = "UNKNOWN"
            record["rainfall_duration_text"] = (
                "Cited article does not report duration; workbook note was: "
                + record["source_duration_raw"]
            )

    # Preserve same-place/same-date reports and mark them for human review.
    overlap_groups: dict[tuple[str, str, str], list[dict]] = {}
    for record in records:
        key = (record["city"].casefold(), record["event_date_raw"].casefold(), (record["location_name"] or "").casefold())
        overlap_groups.setdefault(key, []).append(record)
        if "conflict" in record["rainfall_scope_period"].casefold() or "differ" in record["rainfall_scope_period"].casefold():
            record["review_flags"].append("CONFLICTING_PUBLISHED_RAINFALL_CONTEXT")
    for group in overlap_groups.values():
        if len(group) > 1 and len({record["source_url"] for record in group}) > 1:
            for record in group:
                record["review_flags"].append("OVERLAPPING_SOURCE_REPORTS_SAME_LOCATION_DATE")

    workbook_summary = {
        "sheet_names": book.sheetnames,
        "incident_headers": headers,
        "incident_records": len(records),
        "rows_by_city": {city: sum(record["city"] == city for record in records) for city in sorted({r["city"] for r in records})},
        "event_groups": len({record["event_id"] for record in records}),
        "distinct_location_names": len({(record["city"], (record["location_name"] or "").casefold()) for record in records}),
        "rainfall_amount_records": sum(record["rainfall_amount_mm"] is not None for record in records),
        "duration_records": sum(record["rainfall_duration_minutes"] is not None for record in records),
        "overlap_groups": sum(len(group) > 1 for group in overlap_groups.values()),
    }
    book.close()
    return records, workbook_summary


def csv_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, list):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    if not args.workbook.is_file():
        raise FileNotFoundError(f"Source workbook not found: {args.workbook}")

    records, summary = load_records(args.workbook)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    json_path = args.output_dir / "historical_waterlogging.json"
    csv_path = args.output_dir / "historical_waterlogging.csv"

    with json_path.open("w", encoding="utf-8", newline="\n") as stream:
        json.dump(records, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write("\n")
    with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=FIELDS, extrasaction="raise")
        writer.writeheader()
        writer.writerows({key: csv_value(value) for key, value in record.items()} for record in records)

    print(json.dumps({"workbook": str(args.workbook), "outputs": [str(csv_path), str(json_path)], **summary}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
